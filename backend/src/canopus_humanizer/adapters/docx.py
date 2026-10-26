"""DOCX format adapter with ZIP-inspection and package round-trip validation."""

import io
from typing import Any, List, Optional, Tuple, Union

try:
    import docx
except ImportError:
    docx = None

from canopus_humanizer.adapters.base import BaseAdapter
from canopus_humanizer.engine.errors import FormatError, PARSER_FAILED, SERIALIZATION_FAILED
from canopus_humanizer.engine.models import DocumentBlock, DocumentModel, TransformRequest, TransformResult
from canopus_humanizer.security.archive import SafeArchiveHandler


class DocxAdapter(BaseAdapter):
    """Word DOCX adapter validating package boundaries, table cells, and runs."""

    def __init__(self, archive_handler: Optional[SafeArchiveHandler] = None):
        self.archive_handler = archive_handler or SafeArchiveHandler()

    @property
    def format_name(self) -> str:
        return "docx"

    def parse(self, raw_input: bytes) -> DocumentModel:
        """Validates DOCX package safety and parses text blocks from paragraphs and tables."""
        if docx is None:
            raise FormatError("python-docx is not installed in runtime environment", code="DEPENDENCY_MISSING")

        if not isinstance(raw_input, bytes):
            raise FormatError("DOCX raw_input must be bytes", code="INPUT_INVALID")

        # Validate ZIP bomb and path traversal safety first
        self.archive_handler.validate_archive_bytes(raw_input)

        try:
            doc = docx.Document(io.BytesIO(raw_input))
        except Exception as e:
            raise FormatError(f"Failed to parse DOCX package: {e}", code=PARSER_FAILED)

        blocks: List[DocumentBlock] = []
        idx = 0
        offset = 0

        # 1. Main body paragraphs
        for p in doc.paragraphs:
            text = p.text
            if text.strip():
                blocks.append(
                    DocumentBlock(
                        block_id=f"p_{idx:04d}",
                        block_type="paragraph",
                        start=offset,
                        end=offset + len(text),
                        content=text,
                    )
                )
                idx += 1
                offset += len(text) + 1

        # 2. Table cells
        for t_idx, table in enumerate(doc.tables):
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    for p in cell.paragraphs:
                        text = p.text
                        if text.strip():
                            blocks.append(
                                DocumentBlock(
                                    block_id=f"cell_{t_idx}_{r_idx}_{c_idx}_{idx:04d}",
                                    block_type="table_cell",
                                    start=offset,
                                    end=offset + len(text),
                                    content=text,
                                    metadata={"table": t_idx, "row": r_idx, "col": c_idx},
                                )
                            )
                            idx += 1
                            offset += len(text) + 1

        return DocumentModel(
            blocks=tuple(blocks),
            format_name=self.format_name,
            source_metadata={"block_count": len(blocks)},
        )

    def transform(
        self,
        raw_input: bytes,
        pipeline: Any,
        request: Optional[TransformRequest] = None,
        **kwargs,
    ) -> Tuple[bytes, TransformResult]:
        """Transforms DOCX document and round-trip validates package integrity."""
        if docx is None:
            raise FormatError("python-docx is not installed in runtime environment", code="DEPENDENCY_MISSING")

        self.archive_handler.validate_archive_bytes(raw_input)

        try:
            doc = docx.Document(io.BytesIO(raw_input))
        except Exception as e:
            raise FormatError(f"Failed to open DOCX document: {e}", code=PARSER_FAILED)

        all_changes = []
        all_issues = []
        combined_original = []
        combined_transformed = []

        def process_paragraph(para):
            orig_text = para.text
            if not orig_text.strip():
                return
            combined_original.append(orig_text)

            req = TransformRequest(
                text=orig_text,
                mode=request.mode if request else "conservative",
                ruleset_ids=request.ruleset_ids if request else ("base",),
            )
            res = pipeline.process(req)
            all_changes.extend(res.changes)
            all_issues.extend(res.issues)
            combined_transformed.append(res.transformed_text)

            if res.transformed_text != orig_text:
                if para.runs:
                    # Update text while preserving primary formatting run
                    para.runs[0].text = res.transformed_text
                    for r in para.runs[1:]:
                        r.text = ""
                else:
                    para.text = res.transformed_text

        # Process paragraphs
        for para in doc.paragraphs:
            process_paragraph(para)

        # Process table cells
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        process_paragraph(para)

        # Serialize to in-memory bytes
        out_buf = io.BytesIO()
        try:
            doc.save(out_buf)
            out_bytes = out_buf.getvalue()
        except Exception as e:
            raise FormatError(f"Failed to serialize DOCX document: {e}", code=SERIALIZATION_FAILED)

        # Round-trip integrity check: reopen generated document
        try:
            reopened = docx.Document(io.BytesIO(out_bytes))
            # Verify body paragraph count matches
            if len(reopened.paragraphs) != len(doc.paragraphs):
                raise FormatError("Round-trip validation error: paragraph count discrepancy", code=SERIALIZATION_FAILED)
        except Exception as e:
            raise FormatError(f"Reopened DOCX package validation failed: {e}", code=SERIALIZATION_FAILED)

        orig_full = "\n\n".join(combined_original)
        trans_full = "\n\n".join(combined_transformed)

        result = TransformResult(
            original_text=orig_full,
            transformed_text=trans_full,
            changes=tuple(all_changes),
            issues=tuple(all_issues),
            engine_version=pipeline.engine_version,
            ruleset_versions=pipeline.registry.get_version_info(
                request.ruleset_ids if request else ("base",)
            ),
            accepted=not any(i.severity == "error" for i in all_issues),
        )

        return out_bytes, result
