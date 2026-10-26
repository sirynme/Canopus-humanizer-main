"""End-to-end integration tests for multi-format pipeline processing."""

import io
import docx
import pytest

from canopus_humanizer.adapters.code_comments import CodeCommentAdapter
from canopus_humanizer.adapters.docx import DocxAdapter
from canopus_humanizer.adapters.latex import LatexAdapter
from canopus_humanizer.adapters.plain_text import PlainTextAdapter
from canopus_humanizer.engine.models import TransformRequest
from canopus_humanizer.engine.pipeline import Pipeline
from canopus_humanizer.rules.registry import RuleRegistry


@pytest.fixture
def e2e_pipeline():
    reg = RuleRegistry()
    return Pipeline(registry=reg, version="1.0.0")


def test_e2e_multi_paragraph_plain_text(e2e_pipeline):
    text = (
        "Due to the fact that prior studies were conducted, we delve into this research paper.\n\n"
        "In order to evaluate the benchmark, 1,250 runs achieved 98.6% with citations [1, 2].\n\n"
        "In conclusion, the results are consistent with the hypothesis."
    )
    req = TransformRequest(text=text, ruleset_ids=("base", "academic"))
    result = e2e_pipeline.process(req)

    assert result.accepted is True
    assert len(result.changes) >= 4
    assert result.diff != ""

    # Verify original preserved
    assert result.original_text == text

    # Verify transformations
    trans = result.transformed_text
    assert "Because prior studies were conducted" in trans
    assert "examine closely the present work" in trans
    assert "To evaluate the benchmark" in trans
    assert "1,250" in trans
    assert "98.6%" in trans
    assert "[1, 2]" in trans
    assert "In summary" in trans


def test_e2e_docx_transformation(e2e_pipeline):
    doc = docx.Document()
    doc.add_paragraph("Due to the fact that tests pass, we proceed.")
    buf = io.BytesIO()
    doc.save(buf)
    raw_docx = buf.getvalue()

    adapter = DocxAdapter()
    req = TransformRequest(text="", ruleset_ids=("base",))
    out_docx, result = adapter.transform(raw_docx, pipeline=e2e_pipeline, request=req)

    assert result.accepted is True
    reopened = docx.Document(io.BytesIO(out_docx))
    assert "Because tests pass" in reopened.paragraphs[0].text


def test_e2e_latex_transformation(e2e_pipeline):
    tex = (
        r"\documentclass{article}" "\n"
        r"\begin{document}" "\n"
        r"Due to the fact that $x \in \mathbb{R}$, we proceed." "\n"
        r"\end{document}"
    )
    adapter = LatexAdapter()
    req = TransformRequest(text="", ruleset_ids=("base",))
    out_tex, result = adapter.transform(tex, pipeline=e2e_pipeline, request=req)

    assert result.accepted is True
    assert r"$x \in \mathbb{R}$" in out_tex
    assert "Because" in out_tex


def test_e2e_code_comment_transformation(e2e_pipeline):
    code = (
        "def process():\n"
        "    # In order to compute the sum\n"
        "    total = 0\n"
        "    return total\n"
    )
    adapter = CodeCommentAdapter()
    req = TransformRequest(text="", ruleset_ids=("base",))
    out_code, result = adapter.transform(code, pipeline=e2e_pipeline, language="python", request=req)

    assert result.accepted is True
    assert "# To compute the sum" in out_code
    assert "total = 0" in out_code
    assert "return total" in out_code
