"""Unit tests for Word DOCX format adapter."""

import io
import docx
import pytest
from canopus_humanizer.adapters.docx import DocxAdapter
from canopus_humanizer.engine.models import TransformRequest
from canopus_humanizer.engine.pipeline import Pipeline
from canopus_humanizer.rules.registry import RuleRegistry


def create_sample_docx_bytes():
    """Generates valid DOCX document in memory."""
    doc = docx.Document()
    doc.add_paragraph("Due to the fact that tests are run, we succeed.")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].paragraphs[0].text = "In order to verify, we test tables."
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_adapter_parse_and_transform():
    raw_docx = create_sample_docx_bytes()
    adapter = DocxAdapter()
    model = adapter.parse(raw_docx)
    assert len(model.blocks) >= 2

    registry = RuleRegistry()
    pipeline = Pipeline(registry=registry)
    req = TransformRequest(text="", ruleset_ids=("base",))

    out_bytes, res = adapter.transform(raw_docx, pipeline=pipeline, request=req)
    assert isinstance(out_bytes, bytes)
    assert res.accepted is True
    assert len(res.changes) >= 2

    # Verify output can be cleanly reopened by python-docx
    reopened = docx.Document(io.BytesIO(out_bytes))
    full_text = "\n".join(p.text for p in reopened.paragraphs)
    assert "Because tests are run" in full_text or "because tests are run" in full_text
